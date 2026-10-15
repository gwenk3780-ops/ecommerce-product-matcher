"""
TK热销品 × 站内同品匹配 — Streamlit 版
文本:图像 = 1:9，固定权重
"""

import os, hashlib, warnings, time, json, io
import numpy as np
import pandas as pd
import requests
import torch
import streamlit as st
from PIL import Image as PILImage
from io import BytesIO
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity
from transformers import CLIPProcessor, CLIPModel
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

warnings.filterwarnings('ignore')

# ── 配置（修改这里来调整路径和参数）─────────────────────────────────────────
TEXT_WEIGHT  = 0.1
IMAGE_WEIGHT = 0.9
TOP_N        = 3
_BASE = os.path.dirname(os.path.abspath(__file__))
_cfg_path = os.path.join(_BASE, 'config.json')
if os.path.exists(_cfg_path):
    with open(_cfg_path, encoding='utf-8') as _f:
        _cfg = json.load(_f)
    CACHE_DIR     = os.path.expanduser(_cfg.get('cache_dir',     os.path.join(_BASE, 'image_cache')))
    INDEX_FILE    = os.path.expanduser(_cfg.get('index_file',    os.path.join(_BASE, 'internal_embeddings.npz')))
    INTERNAL_FILE = os.path.expanduser(_cfg.get('internal_file', os.path.join(_BASE, '站内男装总表.xlsx')))
else:
    CACHE_DIR     = os.path.join(_BASE, 'image_cache')
    INDEX_FILE    = os.path.join(_BASE, 'internal_embeddings.npz')
    INTERNAL_FILE = os.path.join(_BASE, '站内男装总表.xlsx')

CATEGORY_MAP = {
    '上衣': ['上衣', '帽衫&运动衫', '毛衫& 开衫', '夹克,外套&马甲', '西装', '套装', '职业装', '传统服饰'],
    '下装': ['短裤', '裤装', '牛仔裤'],
    '内衣': ['贴身衣物&内衣', '睡衣', '袜子'],
}

MATCH_COLORS = ['#E2EFDA', '#FFF2CC', '#FCE4D6']

os.makedirs(CACHE_DIR, exist_ok=True)

# ── 页面配置 ──────────────────────────────────────────────────────────────────
st.set_page_config(
    page_title='TK同品匹配',
    page_icon='🔍',
    layout='wide',
)

st.title('🔍 TK热销品 · 站内同品匹配')

# ── 加载共享资源（缓存，只加载一次）──────────────────────────────────────────
@st.cache_resource(show_spinner='加载图像索引...')
def load_index():
    if not os.path.exists(INDEX_FILE):
        st.error(f'找不到索引文件 {INDEX_FILE}，请先运行 precompute_embeddings.py')
        st.stop()
    idx       = np.load(INDEX_FILE, allow_pickle=True)
    embeds    = idx['embeddings'].astype(np.float32)   # (N, 512)
    urls      = idx['urls'].tolist()
    url_to_idx = {u: i for i, u in enumerate(urls)}   # URL → 行号（轻量）
    return embeds, url_to_idx, urls

@st.cache_resource(show_spinner='加载 CLIP 模型（首次较慢）...')
def load_clip():
    device    = 'mps' if torch.backends.mps.is_available() else 'cpu'
    model     = CLIPModel.from_pretrained('openai/clip-vit-base-patch32').to(device)
    processor = CLIPProcessor.from_pretrained('openai/clip-vit-base-patch32')
    model.eval()
    return model, processor, device

@st.cache_data(show_spinner='读取站内商品数据...')
def load_internal():
    df = pd.read_excel(INTERNAL_FILE, sheet_name='Sheet1')
    df['item_id'] = df['item_id'].fillna('').astype(str)
    for col in ['item_name','merchant_name','sku_price','item_url',
                'second_category_name','active_status','first_image_url']:
        df[col] = df[col].fillna('').astype(str).str.strip()
    return df

idx_embeds, url_to_idx, idx_urls = load_index()
internal   = load_internal()

# ── 工具函数 ──────────────────────────────────────────────────────────────────
def cache_path(url):
    return os.path.join(CACHE_DIR, hashlib.md5(url.encode()).hexdigest() + '.jpg')

def download_image(url):
    if not url or url in ('', 'nan', 'None'):
        return None
    p = cache_path(url)
    if os.path.exists(p):
        try: return PILImage.open(p).convert('RGB')
        except: pass
    try:
        r = requests.get(url, timeout=8, headers={'User-Agent': 'Mozilla/5.0'})
        img = PILImage.open(BytesIO(r.content)).convert('RGB')
        img.save(p, quality=85)
        return img
    except: return None

def encode_tk_images(urls, progress_cb=None):
    model, processor, dev = load_clip()
    result = {}
    todo = [u for u in urls if u and u not in url_to_idx]
    for i in range(0, len(todo), 32):
        batch = todo[i:i+32]
        imgs, valid = [], []
        for u in batch:
            img = download_image(u)
            if img: imgs.append(img); valid.append(u)
        if imgs:
            inp = processor(images=imgs, return_tensors='pt').to(dev)
            with torch.no_grad():
                feats = model.get_image_features(**inp)
                if not isinstance(feats, torch.Tensor):
                    feats = feats.image_embeds if hasattr(feats, 'image_embeds') else feats.pooler_output
                feats = feats / feats.norm(dim=-1, keepdim=True)
            for u, v in zip(valid, feats.cpu().float().numpy()):
                result[u] = v
        if progress_cb:
            progress_cb(min(i+32, len(todo)), len(todo))
    for u in urls:
        if u not in result:
            idx_i = url_to_idx.get(u)
            result[u] = idx_embeds[idx_i] if idx_i is not None else None
    return result

def _render_preview(preview_data):
    import streamlit.components.v1 as components
    data_json = json.dumps(preview_data, ensure_ascii=False)
    html = f"""<!DOCTYPE html><html lang="zh"><head><meta charset="UTF-8">
<style>
*{{box-sizing:border-box;margin:0;padding:0}}
body{{font-family:Arial,sans-serif;font-size:13px;background:#f0f4f8}}
.tabs{{display:flex;background:#16375a;padding:0 12px}}
.tab{{padding:10px 18px;cursor:pointer;color:#aac;font-size:13px;border-bottom:3px solid transparent;transition:.15s}}
.tab.active{{color:#fff;border-bottom-color:#5bc8ff}}
.tab:hover:not(.active){{color:#ddd}}
.content{{padding:12px 16px}}
.stats{{display:flex;gap:10px;margin-bottom:12px;flex-wrap:wrap}}
.stat-card{{background:#fff;border-radius:8px;padding:10px 16px;box-shadow:0 1px 4px rgba(0,0,0,.08)}}
.stat-card .num{{font-size:20px;font-weight:700;color:#1F4E79}}
.stat-card .label{{font-size:11px;color:#888;margin-top:2px}}
.search-bar{{margin-bottom:10px}}
.search-bar input{{padding:6px 12px;border:1px solid #ccc;border-radius:6px;width:300px;font-size:13px;outline:none}}
.search-bar input:focus{{border-color:#1F4E79}}
.table-wrap{{overflow-x:auto;border-radius:8px;box-shadow:0 1px 6px rgba(0,0,0,.1)}}
table{{border-collapse:collapse;width:100%;background:#fff;table-layout:fixed}}
thead tr{{background:#1F4E79;color:#fff}}
thead th{{padding:8px 10px;text-align:left;font-size:12px;white-space:nowrap;position:sticky;top:0;z-index:1}}
tbody tr:nth-child(even) td.base{{background:#EBF3FB}}
tbody tr:hover td{{filter:brightness(.96)}}
td{{padding:7px 6px;border-bottom:1px solid #e8e8e8;vertical-align:top;overflow:hidden}}
td.name-cell{{font-size:12px;line-height:1.5;word-break:break-all}}
.badge{{display:inline-block;padding:2px 7px;border-radius:10px;font-size:11px;font-weight:600}}
.badge-yes{{background:#d4edda;color:#155724}}.badge-no{{background:#f8d7da;color:#721c24}}
.match-group{{margin-bottom:5px;padding:5px 7px;border-radius:6px;display:flex;gap:7px;align-items:flex-start}}
.m1{{background:#E2EFDA}}.m2{{background:#FFF2CC}}.m3{{background:#FCE4D6}}
.match-rank{{font-size:10px;font-weight:700;color:#555;margin-bottom:2px}}
.match-name{{font-size:12px;line-height:1.4;color:#111}}
.match-meta{{font-size:11px;color:#555;margin-top:3px}}
.match-meta a{{color:#1F4E79;text-decoration:none}}
.match-meta a:hover{{text-decoration:underline}}
.sim-bar{{height:4px;background:#e0e0e0;border-radius:2px;margin-top:4px}}
.sim-fill{{height:4px;border-radius:2px}}
.sim-high .sim-fill{{background:#2d8a4e}}.sim-mid .sim-fill{{background:#f0a500}}.sim-low .sim-fill{{background:#c9460a}}
.status-up{{background:#d4edda;color:#155724}}.status-down{{background:#f8d7da;color:#721c24}}.status-other{{background:#e2e3e5;color:#383d41}}
.no-match{{color:#aaa;font-size:12px;font-style:italic}}
.img-thumb{{width:48px;height:48px;object-fit:cover;border-radius:4px;cursor:pointer;flex-shrink:0}}
#lightbox{{display:none;position:fixed;inset:0;z-index:9999;background:rgba(0,0,0,.82);align-items:center;justify-content:center;cursor:zoom-out}}
#lightbox.open{{display:flex}}
#lightbox img{{max-width:88vw;max-height:88vh;border-radius:8px;box-shadow:0 8px 40px rgba(0,0,0,.6);cursor:default;object-fit:contain}}
#lightbox-close{{position:fixed;top:14px;right:20px;color:#fff;font-size:28px;cursor:pointer;opacity:.8;user-select:none}}
#lightbox-close:hover{{opacity:1}}
</style></head><body>
<div id="lightbox" onclick="closeLightbox()"><span id="lightbox-close" onclick="closeLightbox()">✕</span><img id="lightbox-img" src="" onclick="event.stopPropagation()"></div>
<div class="tabs" id="tabs"></div>
<div class="content">
  <div class="stats" id="stats"></div>
  <div class="search-bar"><input type="text" id="search" placeholder="搜索商品名称…" oninput="filterTable()"></div>
  <div class="table-wrap"><table>
    <thead><tr>
      <th style="width:46px">图片</th>
      <th style="width:160px">TK商品名称</th>
      <th style="width:70px">TK店铺</th>
      <th style="width:55px">TK价格</th>
      <th style="width:50px">站内同品</th>
      <th style="width:190px">同品1 <span style="font-weight:400;color:#adf">(最高相似)</span></th>
      <th style="width:190px">同品2</th>
      <th style="width:190px">同品3</th>
    </tr></thead>
    <tbody id="tbody"></tbody>
  </table></div>
</div>
<script>
const allData = {data_json};
let currentTab = Object.keys(allData)[0];
(function init(){{
  const tabsEl = document.getElementById('tabs');
  Object.keys(allData).forEach(t => {{
    const d = document.createElement('div');
    d.className = 'tab' + (t===currentTab?' active':'');
    d.textContent = t; d.onclick = () => switchTab(t);
    tabsEl.appendChild(d);
  }});
  renderTable();
}})();
function switchTab(t){{
  currentTab=t;
  document.querySelectorAll('.tab').forEach(el=>el.classList.toggle('active',el.textContent===t));
  document.getElementById('search').value='';
  renderTable();
}}
function simClass(s){{ return s>=0.7?'sim-high':s>=0.4?'sim-mid':'sim-low'; }}
function statusBadge(s){{
  const cls=s==='上架'?'status-up':(s&&s.includes('下架')?'status-down':'status-other');
  return `<span class="badge ${{cls}}">${{s||'-'}}</span>`;
}}
function matchBlock(r,rank,cc){{
  const name=r[`同品${{rank}}_名称`]||'';
  if(!name) return '<div class="no-match">—</div>';
  const img=r[`同品${{rank}}_图片`]||'',shop=r[`同品${{rank}}_店铺`]||'',price=r[`同品${{rank}}_价格`]||'';
  const sim=parseFloat(r[`同品${{rank}}_相似度`])||0,status=r[`同品${{rank}}_状态`]||'',url=r[`同品${{rank}}_链接`]||'';
  const pct=Math.round(sim*100);
  return `<div class="match-group ${{cc}}">
    ${{img?`<img class="img-thumb" src="${{img}}" referrerpolicy="no-referrer" loading="lazy" onclick="openLightbox(this.src)">`:''}}
    <div style="flex:1;min-width:0">
      <div class="match-rank">#${{rank}} 相似度 ${{(sim*100).toFixed(1)}}%</div>
      <div class="match-name">${{name}}</div>
      <div class="match-meta">🏪 ${{shop}} &nbsp;|&nbsp; 💰 ${{price}} &nbsp;|&nbsp; ${{statusBadge(status)}}${{url?` &nbsp;| <a href="${{url}}" target="_blank">详情页↗</a>`:''}}</div>
      <div class="sim-bar ${{simClass(sim)}}"><div class="sim-fill" style="width:${{pct}}%"></div></div>
    </div></div>`;
}}
function renderTable(filter=''){{
  const rows=(allData[currentTab]||[]).filter(r=>{{
    if(!filter) return true;
    const q=filter.toLowerCase();
    return (r['TK商品名称']||'').toLowerCase().includes(q)||(r['TK店铺']||'').toLowerCase().includes(q)||(r['同品1_名称']||'').toLowerCase().includes(q);
  }});
  const all=allData[currentTab]||[];
  const high=all.filter(r=>(parseFloat(r['同品1_相似度'])||0)>=0.7).length;
  const mid=all.filter(r=>{{const s=parseFloat(r['同品1_相似度'])||0;return s>=0.4&&s<0.7;}}).length;
  document.getElementById('stats').innerHTML=`
    <div class="stat-card"><div class="num">${{all.length}}</div><div class="label">TK热销品</div></div>
    <div class="stat-card"><div class="num" style="color:#2d8a4e">${{high}}</div><div class="label">高相似 ≥70%</div></div>
    <div class="stat-card"><div class="num" style="color:#f0a500">${{mid}}</div><div class="label">中相似 40-70%</div></div>
    <div class="stat-card"><div class="num" style="color:#c9460a">${{all.length-high-mid}}</div><div class="label">低相似 &lt;40%</div></div>
    <div class="stat-card"><div class="num">${{rows.length}}</div><div class="label">当前显示</div></div>`;
  document.getElementById('tbody').innerHTML=rows.map(r=>{{
    const img=r['TK图片']||'',hs=r['站内是否有同品'];
    const sb=hs==='是'?'<span class="badge badge-yes">是</span>':hs==='否'?'<span class="badge badge-no">否</span>':(hs||'');
    return `<tr>
      <td class="base">${{img?`<img class="img-thumb" src="${{img}}" referrerpolicy="no-referrer" loading="lazy" onclick="openLightbox(this.src)">`:''}}  </td>
      <td class="base name-cell">${{r['TK商品名称']||''}}</td>
      <td class="base">${{r['TK店铺']||''}}</td>
      <td class="base">${{r['TK价格']||''}}</td>
      <td class="base" style="text-align:center">${{sb}}</td>
      <td>${{matchBlock(r,1,'m1')}}</td>
      <td>${{matchBlock(r,2,'m2')}}</td>
      <td>${{matchBlock(r,3,'m3')}}</td>
    </tr>`;
  }}).join('');
}}
function filterTable(){{ renderTable(document.getElementById('search').value); }}
function openLightbox(src){{ document.getElementById('lightbox-img').src=src; document.getElementById('lightbox').classList.add('open'); }}
function closeLightbox(){{ document.getElementById('lightbox').classList.remove('open'); document.getElementById('lightbox-img').src=''; }}
document.addEventListener('keydown',e=>{{ if(e.key==='Escape') closeLightbox(); }});
</script></body></html>"""
    components.html(html, height=900, scrolling=True)

def build_excel(result_df):
    wb = Workbook()
    wb.remove(wb.active)
    HEADER_FILL = PatternFill('solid', start_color='1F4E79')
    HEADER_FONT = Font(name='Arial', bold=True, color='FFFFFF', size=10)
    DATA_FONT   = Font(name='Arial', size=10)
    ALT_FILL    = PatternFill('solid', start_color='EBF3FB')
    MFILLS = [PatternFill('solid', start_color=c.lstrip('#')) for c in ['E2EFDA','FFF2CC','FCE4D6']]

    col_order = ['TK分类','TK商品名称','TK图片','TK价格','TK店铺','站内是否有同品']
    for i in range(1, TOP_N+1):
        col_order += [f'同品{i}_id',f'同品{i}_名称',f'同品{i}_图片',f'同品{i}_店铺',
                      f'同品{i}_价格',f'同品{i}_相似度',f'同品{i}_状态',f'同品{i}_链接']
    widths = {'TK分类':8,'TK商品名称':45,'TK图片':15,'TK价格':10,'TK店铺':20,'站内是否有同品':12}
    for i in range(1,TOP_N+1):
        widths.update({f'同品{i}_id':15,f'同品{i}_名称':40,f'同品{i}_图片':15,f'同品{i}_店铺':20,
                       f'同品{i}_价格':10,f'同品{i}_相似度':8,f'同品{i}_状态':10,f'同品{i}_链接':20})

    for cat in CATEGORY_MAP:
        sub = result_df[result_df['TK分类']==cat].copy()
        if sub.empty: continue
        ws = wb.create_sheet(title=cat)
        headers = [c for c in col_order if c in sub.columns]
        for ci, col in enumerate(headers, 1):
            cell = ws.cell(row=1, column=ci, value=col)
            cell.font = HEADER_FONT; cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        for ri, (_, row) in enumerate(sub.iterrows(), 2):
            for ci, col in enumerate(headers, 1):
                val = row.get(col, '')
                if pd.isna(val): val = ''
                cell = ws.cell(row=ri, column=ci, value=val)
                cell.font = DATA_FONT
                cell.alignment = Alignment(vertical='center', wrap_text=True)
                for rank in range(1, TOP_N+1):
                    if col.startswith(f'同品{rank}_'):
                        cell.fill = MFILLS[rank-1]; break
                else:
                    if ri % 2 == 0: cell.fill = ALT_FILL
        for ci, col in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(ci)].width = widths.get(col, 15)
        ws.row_dimensions[1].height = 28
        ws.freeze_panes = 'A2'

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

# ── 主界面 ────────────────────────────────────────────────────────────────────
uploaded = st.file_uploader(
    '上传 TK 商品 Excel 文件（需包含上衣/下装/内衣三个 Sheet）',
    type=['xlsx'],
)

if uploaded:
    t0 = time.time()

    # 读取 TK 文件
    tk_sheets = pd.read_excel(uploaded, sheet_name=None)
    tk_rows = []
    for sheet_name, df in tk_sheets.items():
        if sheet_name not in CATEGORY_MAP: continue
        img_col = '商品图片链接' if '商品图片链接' in df.columns else '商品图片'
        for _, row in df.iterrows():
            tk_rows.append({
                'TK分类':     sheet_name,
                'TK商品名称': str(row.get('商品名称', '')).strip(),
                'TK图片':     str(row.get(img_col, '')).strip(),
                'TK价格':     str(row.get('商品售价', '')).strip(),
                'TK店铺':     str(row.get('店铺名', '')).strip(),
                '站内是否有同品': str(row.get('站内是否有同品', '')).strip(),
            })
    tk_df = pd.DataFrame(tk_rows)
    st.success(f'已读取 {len(tk_df)} 个TK商品（上衣/下装/内衣）')

    if st.button('🚀 开始匹配', type='primary', use_container_width=True):
        results = []
        status  = st.status('匹配中...', expanded=True)

        with status:
            for category, internal_cats in CATEGORY_MAP.items():
                tk_sub = tk_df[tk_df['TK分类'] == category].copy()
                if tk_sub.empty: continue
                int_sub = internal[internal['second_category_name'].isin(internal_cats)].copy()
                st.write(f'**[{category}]** TK: {len(tk_sub)} 个，站内候选: {len(int_sub):,} 个')

                if int_sub.empty:
                    for _, row in tk_sub.iterrows():
                        r = row.to_dict()
                        for i in range(1, TOP_N+1):
                            for k in ['名称','图片','店铺','价格','链接','状态','相似度']:
                                r[f'同品{i}_{k}'] = ''
                        results.append(r)
                    continue

                int_sub_r = int_sub.reset_index(drop=True)

                # TF-IDF
                all_names    = list(tk_sub['TK商品名称']) + list(int_sub['item_name'])
                vec          = TfidfVectorizer(analyzer='char_wb', ngram_range=(2,4), min_df=1, sublinear_tf=True)
                mat          = vec.fit_transform([n.lower() for n in all_names])
                text_sims    = cosine_similarity(mat[:len(tk_sub)], mat[len(tk_sub):])

                # TK图像向量
                prog = st.progress(0, text=f'下载 {category} TK商品图片...')
                def prog_cb(done, total):
                    prog.progress(done/total, text=f'下载图片 {done}/{total}')
                tk_embeds = encode_tk_images(tk_sub['TK图片'].tolist(), prog_cb)
                prog.empty()

                # 站内图像矩阵
                int_urls = int_sub_r['first_image_url'].tolist()
                int_idx_list = [url_to_idx.get(u) for u in int_urls]
                int_vecs_img = np.array(
                    [idx_embeds[i] if i is not None else np.zeros(512, dtype=np.float32)
                     for i in int_idx_list],
                    dtype=np.float32
                )
                has_img = np.array([i is not None for i in int_idx_list])

                for idx, (_, row) in enumerate(tk_sub.iterrows()):
                    t_sims   = text_sims[idx]
                    tk_embed = tk_embeds.get(row['TK图片'])

                    if tk_embed is not None:
                        img_sims = (int_vecs_img @ tk_embed + 1) / 2
                        img_sims[~has_img] = 0.0
                        combined = np.where(has_img,
                            TEXT_WEIGHT*t_sims + IMAGE_WEIGHT*img_sims, t_sims)
                    else:
                        combined = t_sims

                    top_idx = np.argsort(combined)[::-1][:TOP_N]
                    r = row.to_dict()
                    for rank, i in enumerate(top_idx, 1):
                        m = int_sub_r.iloc[i]
                        r[f'同品{rank}_id']     = m['item_id']
                        r[f'同品{rank}_名称']   = m['item_name']
                        r[f'同品{rank}_图片']   = m['first_image_url']
                        r[f'同品{rank}_店铺']   = m['merchant_name']
                        r[f'同品{rank}_价格']   = m['sku_price']
                        r[f'同品{rank}_链接']   = m['item_url']
                        r[f'同品{rank}_状态']   = m['active_status']
                        r[f'同品{rank}_相似度'] = round(float(combined[i]), 4)
                    results.append(r)

            status.update(label='✅ 匹配完成！', state='complete')

        result_df = pd.DataFrame(results)
        elapsed   = time.time() - t0
        st.info(f'共处理 {len(result_df)} 个TK商品，耗时 {elapsed:.1f} 秒')

        # 下载按钮
        excel_bytes = build_excel(result_df)
        st.download_button(
            label='⬇️ 下载结果 Excel',
            data=excel_bytes,
            file_name=f'TK同品匹配结果.xlsx',
            mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            type='primary',
            use_container_width=True,
        )

        # 预览
        st.subheader('结果预览')
        preview_data = {}
        for cat in CATEGORY_MAP:
            sub = result_df[result_df['TK分类']==cat].copy().fillna('')
            preview_data[cat] = sub.to_dict(orient='records')
        _render_preview(preview_data)
