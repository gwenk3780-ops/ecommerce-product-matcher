"""
TK热销品 × 站内同品匹配（纯文本版，无需图片）
基于 TF-IDF 字符级相似度，运行时间 < 1 分钟。
"""

import warnings
import json

import numpy as np
import pandas as pd
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

warnings.filterwarnings('ignore')

# ── 配置 ──────────────────────────────────────────────────────────────────────
TOP_N = 3

# ── 1. 读取数据 ──────────────────────────────────────────────────────────────
print("读取文件...")
tk_sheets = pd.read_excel('TK_top300盘点.xlsx', sheet_name=None)
internal  = pd.read_excel('站内男装总表.xlsx', sheet_name='Sheet1')

internal = internal[internal['status'] == '正常'].copy()
internal['item_name']            = internal['item_name'].fillna('').str.strip()
internal['merchant_name']        = internal['merchant_name'].fillna('')
internal['sku_price']            = internal['sku_price'].fillna('')
internal['item_url']             = internal['item_url'].fillna('')
internal['second_category_name'] = internal['second_category_name'].fillna('')
internal['active_status']        = internal['active_status'].fillna('')
internal['first_image_url']      = internal['first_image_url'].fillna('')

print(f"站内商品总数（正常）: {len(internal)}")

# ── 2. 类目映射 ──────────────────────────────────────────────────────────────
CATEGORY_MAP = {
    '上衣': ['上衣', '帽衫&运动衫', '毛衫& 开衫', '夹克,外套&马甲', '西装', '套装', '职业装', '传统服饰'],
    '下装': ['短裤', '裤装', '牛仔裤'],
    '内衣': ['贴身衣物&内衣', '睡衣', '袜子'],
}

# ── 3. 合并TK商品 ────────────────────────────────────────────────────────────
tk_rows = []
for sheet_name, df in tk_sheets.items():
    img_col = '商品图片链接' if '商品图片链接' in df.columns else '商品图片'
    for _, row in df.iterrows():
        tk_rows.append({
            'TK分类':        sheet_name,
            'TK商品名称':    str(row.get('商品名称', '')).strip(),
            'TK图片':        str(row.get(img_col, '')).strip(),
            'TK价格':        str(row.get('商品售价', '')).strip(),
            'TK店铺':        str(row.get('店铺名', '')).strip(),
            '站内是否有同品': str(row.get('站内是否有同品', '')).strip(),
        })

tk_df = pd.DataFrame(tk_rows)
print(f"TK商品总数: {len(tk_df)}")

# ── 4. 按类目匹配 ────────────────────────────────────────────────────────────
results = []

for category, internal_cats in CATEGORY_MAP.items():
    tk_sub       = tk_df[tk_df['TK分类'] == category].copy()
    if tk_sub.empty:
        continue

    internal_sub = internal[internal['second_category_name'].isin(internal_cats)].copy()
    print(f"\n[{category}] TK商品: {len(tk_sub)}, 站内候选: {len(internal_sub)}")

    if internal_sub.empty:
        for _, row in tk_sub.iterrows():
            r = row.to_dict()
            for i in range(1, TOP_N + 1):
                for key in ['名称', '图片', '店铺', '价格', '链接', '状态', '相似度']:
                    r[f'同品{i}_{key}'] = ''
            results.append(r)
        continue

    all_names    = list(tk_sub['TK商品名称']) + list(internal_sub['item_name'])
    vectorizer   = TfidfVectorizer(analyzer='char_wb', ngram_range=(2, 4),
                                   min_df=1, sublinear_tf=True)
    tfidf_matrix = vectorizer.fit_transform([n.lower() for n in all_names])
    tk_vecs      = tfidf_matrix[:len(tk_sub)]
    int_vecs     = tfidf_matrix[len(tk_sub):]
    sim_matrix   = cosine_similarity(tk_vecs, int_vecs)

    internal_sub_reset = internal_sub.reset_index(drop=True)

    for idx, (_, row) in enumerate(tk_sub.iterrows()):
        sims        = sim_matrix[idx]
        top_indices = np.argsort(sims)[::-1][:TOP_N]

        r = row.to_dict()
        for rank, i in enumerate(top_indices, 1):
            match = internal_sub_reset.iloc[i]
            r[f'同品{rank}_名称']   = match['item_name']
            r[f'同品{rank}_图片']   = match['first_image_url']
            r[f'同品{rank}_店铺']   = match['merchant_name']
            r[f'同品{rank}_价格']   = match['sku_price']
            r[f'同品{rank}_链接']   = match['item_url']
            r[f'同品{rank}_状态']   = match['active_status']
            r[f'同品{rank}_相似度'] = round(float(sims[i]), 4)

        for rank in range(len(top_indices) + 1, TOP_N + 1):
            for key in ['名称', '图片', '店铺', '价格', '链接', '状态', '相似度']:
                r[f'同品{rank}_{key}'] = ''

        results.append(r)

result_df = pd.DataFrame(results)

# ── 5. 输出 Excel ────────────────────────────────────────────────────────────
output_path = 'TK热销品站内同品匹配.xlsx'

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

wb = Workbook()
wb.remove(wb.active)

HEADER_FILL = PatternFill('solid', start_color='1F4E79')
HEADER_FONT = Font(name='Arial', bold=True, color='FFFFFF', size=10)
DATA_FONT   = Font(name='Arial', size=10)
ALT_FILL    = PatternFill('solid', start_color='EBF3FB')
MATCH_FILLS = [
    PatternFill('solid', start_color='E2EFDA'),
    PatternFill('solid', start_color='FFF2CC'),
    PatternFill('solid', start_color='FCE4D6'),
]

col_order = ['TK分类', 'TK商品名称', 'TK图片', 'TK价格', 'TK店铺', '站内是否有同品']
for i in range(1, TOP_N + 1):
    col_order += [f'同品{i}_名称', f'同品{i}_图片', f'同品{i}_店铺',
                  f'同品{i}_价格', f'同品{i}_相似度', f'同品{i}_状态', f'同品{i}_链接']

COL_WIDTHS = {'TK分类': 8, 'TK商品名称': 45, 'TK图片': 15,
              'TK价格': 10, 'TK店铺': 20, '站内是否有同品': 12}
for i in range(1, TOP_N + 1):
    COL_WIDTHS.update({f'同品{i}_名称': 40, f'同品{i}_图片': 15,
                       f'同品{i}_店铺': 20, f'同品{i}_价格': 10,
                       f'同品{i}_相似度': 8, f'同品{i}_状态': 10, f'同品{i}_链接': 20})

for category in CATEGORY_MAP.keys():
    sub = result_df[result_df['TK分类'] == category].copy()
    if sub.empty:
        continue

    ws      = wb.create_sheet(title=category)
    headers = [c for c in col_order if c in sub.columns]

    for ci, col in enumerate(headers, 1):
        cell           = ws.cell(row=1, column=ci, value=col)
        cell.font      = HEADER_FONT
        cell.fill      = HEADER_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    for ri, (_, row) in enumerate(sub.iterrows(), 2):
        is_alt = (ri % 2 == 0)
        for ci, col in enumerate(headers, 1):
            val = row.get(col, '')
            if pd.isna(val):
                val = ''
            cell           = ws.cell(row=ri, column=ci, value=val)
            cell.font      = DATA_FONT
            cell.alignment = Alignment(vertical='center', wrap_text=True)
            for rank in range(1, TOP_N + 1):
                if col.startswith(f'同品{rank}_'):
                    cell.fill = MATCH_FILLS[rank - 1]
                    break
            else:
                if is_alt:
                    cell.fill = ALT_FILL

    for ci, col in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(ci)].width = COL_WIDTHS.get(col, 15)
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = 'A2'

wb.save(output_path)
print(f"\n完成！结果已保存到: {output_path}")
print(f"共处理 {len(result_df)} 个TK商品")

# ── 6. 同步生成 preview_data.json ────────────────────────────────────────────
preview_data = {}
for category in CATEGORY_MAP.keys():
    sub = result_df[result_df['TK分类'] == category].copy().fillna('')
    preview_data[category] = sub.to_dict(orient='records')

with open('preview_data.json', 'w', encoding='utf-8') as f:
    json.dump(preview_data, f, ensure_ascii=False)
print("preview_data.json 已更新")
